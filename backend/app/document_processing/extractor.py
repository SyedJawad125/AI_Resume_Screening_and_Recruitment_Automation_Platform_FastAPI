"""
app/document_processing/extractor.py
────────────────────────────────────────
Pipeline:
    1. Detect whether each PDF page has extractable text.
    2. Extract text natively (PyMuPDF) where available.
    3. Fall back to OCR (via the OCR abstraction) for image-only pages.
    4. Normalize whitespace.
    5. Preserve page numbers so evidence can cite "resume page: N".

Also handles DOCX (python-docx has no OCR concept — it's always text).
"""

import io
import re
from dataclasses import dataclass, field

import pymupdf as fitz  # PyMuPDF (the 'fitz' import alias is deprecated upstream)
from PIL import Image

from app.document_processing.ocr import get_ocr_provider

# A page with fewer than this many extractable characters is treated as
# "probably scanned" and sent through OCR instead.
MIN_TEXT_CHARS_PER_PAGE = 20


@dataclass
class ExtractedDocument:
    pages: list[dict] = field(default_factory=list)   # [{"page": 1, "text": "..."}]
    used_ocr: bool = False

    @property
    def full_text(self) -> str:
        return "\n\n".join(p["text"] for p in self.pages)

    @property
    def page_count(self) -> int:
        return len(self.pages)


def _normalize(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pdf(file_bytes: bytes) -> ExtractedDocument:
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    ocr_provider = get_ocr_provider()
    result = ExtractedDocument()

    for page_index, page in enumerate(doc, start=1):
        text = page.get_text().strip()

        if len(text) < MIN_TEXT_CHARS_PER_PAGE:
            # Likely a scanned/image-based page — rasterize and OCR it.
            pixmap = page.get_pixmap(dpi=200)
            image = Image.open(io.BytesIO(pixmap.tobytes("png")))
            text = ocr_provider.extract_text(image)
            result.used_ocr = True

        result.pages.append({"page": page_index, "text": _normalize(text)})

    doc.close()
    return result


def extract_docx(file_bytes: bytes) -> ExtractedDocument:
    import docx

    document = docx.Document(io.BytesIO(file_bytes))
    text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
    return ExtractedDocument(pages=[{"page": 1, "text": _normalize(text)}], used_ocr=False)


def extract_document(file_bytes: bytes, filename: str) -> ExtractedDocument:
    ext = filename.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        return extract_pdf(file_bytes)
    if ext == "docx":
        return extract_docx(file_bytes)
    raise ValueError(f"Unsupported file extension: .{ext}")
